"""Check the generated .docx for layout faults, without rendering them.

Written after a fault the previous check could not see. The builder set cell
widths and the verifier read those same cell widths back, so a table whose grid
said 7.04in — the full page text width, twice the column it sits in — passed.
Word happens to repair that from the cell widths on render; LibreOffice and the
converters journals run do not, so the file was wrong even though Word looked right.

This reads the two things that actually decide the layout: the table grid, and
whether the widest unbreakable token in a cell fits the cell it is in, measured
in the real font rather than estimated.

    python paper/audit_layout.py

Exits non-zero if anything overflows, so it can gate a rebuild.
"""
import os
import sys

import docx
from PIL import ImageFont

EMU_PER_IN = 914400
TWIP_PER_IN = 1440
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = [
    # file, width one line of text may occupy
    ('Explainable_ChestXray_Paper_IJSR.docx', 3.42),          # one IJSR column
    ('Explainable_ChestXray_Paper_singlecolumn.docx', 6.50),  # Letter less 1in margins
]

FONTS = {
    (False, False): r'C:\Windows\Fonts\times.ttf',
    (True, False): r'C:\Windows\Fonts\timesbd.ttf',
    (False, True): r'C:\Windows\Fonts\timesi.ttf',
    (True, True): r'C:\Windows\Fonts\timesbi.ttf',
}
_cache = {}


def width_in(text, bold, italic, pt):
    """Rendered width of a string, in inches, at 96dpi."""
    key = (bool(bold), bool(italic), round(pt, 1))
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FONTS[(bool(bold), bool(italic))],
                                         max(1, int(round(pt * 96 / 72))))
    return _cache[key].getlength(text) / 96.0


def audit(path, line_w):
    doc = docx.Document(path)
    default_pt = doc.styles['Normal'].font.size.pt if doc.styles['Normal'].font.size else 10.0
    problems = []
    print(f'\n{os.path.basename(path)}  ({len(doc.tables)} tables, line width {line_w}in)')

    for ti, t in enumerate(doc.tables, 1):
        grid = [int(g.get(W + 'w')) / TWIP_PER_IN
                for g in t._tbl.tblGrid.findall(W + 'gridCol')]
        total = sum(grid)
        if total > line_w + 0.01:
            problems.append(f'table {ti}: grid is {total:.3f}in against a {line_w}in line')

        # Under an autofit layout the grid is only a starting point — Word widens a
        # column whose content demands it — so a token wider than its column is a
        # note, not a fault. Under a fixed layout the column width is the last word.
        layout = t._tbl.tblPr.find(W + 'tblLayout')
        fixed = layout is not None and layout.get(W + 'type') == 'fixed'

        mar = t._tbl.tblPr.find(W + 'tblCellMar')
        pad = sum((int(e.get(W + 'w')) / TWIP_PER_IN) if e is not None else 0.08
                  for e in (mar.find(W + s) if mar is not None else None
                            for s in ('left', 'right')))

        for ci, col_w in enumerate(grid):
            avail = col_w - pad
            worst = (0.0, '')
            for row in t.rows:
                if ci >= len(row.cells):
                    continue
                for p in row.cells[ci].paragraphs:
                    for run in p.runs:
                        pt = run.font.size.pt if run.font.size else default_pt
                        for tok in run.text.split():
                            w = width_in(tok, run.font.bold, run.font.italic, pt)
                            if w > worst[0]:
                                worst = (w, tok)
            if worst[0] > avail:
                msg = (f'table {ti} col {ci + 1}: {worst[1]!r} needs {worst[0]:.3f}in '
                       f'of {avail:.3f}in')
                if fixed:
                    problems.append(msg + ' — it will break mid-word')
                else:
                    print(f'    note: {msg}; autofit, Word will widen the column')

        print(f'  table {ti}: {len(t.rows)}x{len(t.columns)}, grid {total:.3f}in')

    return problems


if __name__ == '__main__':
    found = []
    for name, line_w in DOCS:
        found += audit(os.path.join(HERE, name), line_w)

    print()
    if found:
        print(f'{len(found)} problem(s):')
        for p in found:
            print('  *', p)
        sys.exit(1)
    print('no layout problems')
